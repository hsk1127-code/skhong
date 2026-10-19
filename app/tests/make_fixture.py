"""실제 사용 중인 프로젝트 엑셀과 같은 구조의 테스트용 통합문서를 만든다.

    python3 app/tests/make_fixture.py <출력 경로.xlsx> [업체 BOM 출력 경로.xlsx] [회의록 출력 경로.xlsx]
"""

import datetime as dt
import io
import sys

import openpyxl
from openpyxl.drawing.image import Image
from PIL import Image as PILImage, ImageDraw


def png(text):
    im = PILImage.new("RGB", (400, 160), "white")
    ImageDraw.Draw(im).text((20, 60), text, fill="black")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return Image(buf)


def main(out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revision"
    ws.append(["NO", "날짜", "작성자", "내용", "비고"])
    ws.append([1, dt.datetime(2026, 10, 5), "홍길동", "신규 작성", None])
    ws["B2"].number_format = 'mm"월"\\ dd"일"'

    ws = wb.create_sheet("개요")
    ws["A1"] = "설비/품목명: 테스트_설비"
    ws.merge_cells("A1:C2")
    for r, row in enumerate([["분류", "항목", "비고"], ["Project", "테스트 AMR 도입"], ["고객사", "테스트고객"],
                             ["사업 범위", "AMR : 2EA"], [None, "Charger : 1EA"]], start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    ws.merge_cells("A6:A7")

    ws = wb.create_sheet("BOM")
    ws["A1"] = "BOM"
    ws.append(["NO", "구분", "품목", "모델명", "업체명", "개수", "비고"])
    ws.append(["지정품 List"])
    ws.merge_cells("A3:G3")
    ws.append(["NO", "구분", "품목", "모델명", "업체명", "개수", "비고"])
    ws.append([1, "제어모듈", "메인제어PC", "Nuvo-7006E", "Neousys", 1, None])
    ws.append([2, "센서", "말굽센서", "PM-R25", "Panasonic", 5, None])

    ws = wb.create_sheet("킥오프 회의록(2026.09.17)")
    ws["B2"] = "회     의     록"
    ws["B5"], ws["E5"], ws["L5"], ws["O5"] = "부 서 명", "영업팀", "작 성자", "홍길동"
    ws["B7"], ws["E7"], ws["L7"], ws["O7"] = "회의일자", "2026.09.16 (수)", "회의장소", "본사 회의실"
    ws["B9"], ws["E9"], ws["L9"], ws["O9"] = "회의시간", "15:00 ~ 18:00", "구분", "  정례  ■ 비정례 "
    ws["B11"], ws["E11"] = "제   목", "고객사 Kick-off"
    ws["B13"], ws["E13"] = "참석자", " - 고객사 : 김담당"
    ws["B18"] = "회 의 내 용"
    ws["B19"], ws["D19"] = "No", "주  요  내  용"
    ws["B20"], ws["D20"] = 1, "개요"
    ws["D21"] = "ⅰ. 회의 목적 : 사양 확정"
    ws["B30"] = "실천 의제"
    ws["B31"], ws["F31"], ws["H31"] = "지 시 내 용", "담 당 자", "기   한"
    ws["B32"], ws["F32"], ws["H32"] = "견적 제출", "영업팀", "미정"
    ws["B34"] = "특이사항"

    # 프로젝트마다 추가되는 문서형 사양서: 번호 제목, 문단, 작은 표
    ws = wb.create_sheet("설계 및 제작 조건")
    ws["B2"] = "1. 설계 및 제작 조건"
    ws["B4"] = "1.1. 환경 조건"
    ws["B6"] = "1) 온도: 22.5 ~ 23.5 °C"
    ws["B7"] = "2) 습도: 45 ± 20% R.H"
    ws["B9"] = "1.2. 업무 Scope"
    for r, row in enumerate([("항목", "내용", "담당"), ("Hardware", "로봇 설계", "외주"), (None, "3D 도면", None)], start=10):
        for c, v in enumerate(row, start=2):
            if v:
                ws.cell(r, c, v)
    ws = wb.create_sheet("전장품 문의(2026.10.06)")
    ws.add_image(png("mail capture"), "A1")

    ws = wb.create_sheet("일정")
    ws["J2"], ws["K2"] = "기준일", dt.datetime(2026, 10, 8)
    for c, v in zip("BDEFGHIJK", ["항목", "세부사항", "비고", "담당자", "시작일", "종료일", "기간", "경과일", "잔여일"]):
        ws[f"{c}3"] = v
    rows = [("사양 협의", "사양협의", "2026-09-11", "2026-09-20"), (None, "- 1차 DR", "2026-10-06", "2026-10-06"),
            ("기구설계", "3D 설계", "2026-09-21", "2099-10-25")]
    for i, (phase, name, s, e) in enumerate(rows, start=4):
        ws[f"C{i}"], ws[f"D{i}"] = phase, name
        ws[f"G{i}"], ws[f"H{i}"] = dt.datetime.fromisoformat(s), dt.datetime.fromisoformat(e)
        ws[f"G{i}"].number_format = ws[f"H{i}"].number_format = "mm-dd-yy"
    ws.merge_cells("C4:C5")
    wb.save(out)


def supplier_bom(out):
    """업체에서 받는 메이커리스트 형식: 제목, 소제목, 영문 머리글, 중복 품목"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MakerList"
    ws["B1"] = "Deployment of a Custom AMR"
    ws.merge_cells("B1:H1")
    ws["B2"] = "전장"
    ws.append([None, "No.", "Item Name", "Spec.", "Maker", "Q'ty", "Remark", "Cost($)"])
    ws.append([None, 1, "Lidar", "Nano Scan3", "SICK", 1, None, 1337.150311639368])
    ws.append([None, 2, "Drive", "Leadshine", "Leadshine", 2, "Driving"])
    ws.append([None, 3, "Lidar", "Nano Scan3", "sick", 1, "spare"])
    ws.append([None, "기구"])
    ws.append([None, 4, "LM Guide", "SSEBL10", "Misumi", 4])
    ws.append([None, 5])
    old = wb.create_sheet("(이전)MakerList(2026.08.24)")
    old.append(["No.", "Item Name", "Spec.", "Maker"])
    old.append([1, "Lidar", "Nano Scan3", "SICK"])
    wb.save(out)


def meeting(out):
    """다른 양식의 회의록: 시트 이름 Sheet1, "일시 : …" 한 칸 표기, 조치사항 표, 사진"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in [["회의록"], ["회의명", "1차 DR 회의"], ["일시 : 2026.10.06 14:00~16:00"], ["장소", "본사 3층", "작성", "홍길동"],
                ["참석자", "고객사 김담당, 홍길동"], [], ["회의 내용", "도면 검토 결과 공유"], [None, "- 수정 사항 3건"], [],
                ["조치사항"], ["No", "내용", "담당", "기한"], [1, "수정 도면 송부", "홍길동", "2026.10.13"]]:
        ws.append(row)
    ws.add_image(png("whiteboard"), "F2")
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1])
    if len(sys.argv) > 2:
        supplier_bom(sys.argv[2])
    if len(sys.argv) > 3:
        meeting(sys.argv[3])
